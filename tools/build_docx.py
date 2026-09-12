"""把兩條路線的部署手冊轉成 Word 文件，內容跟 html 版一致，格式簡化成適合列印/離線閱讀。

用法：在 tools/downloads/ 資料夾底下跑
    python ../build_docx.py .
會在當前資料夾產生／覆蓋兩份 docx。html 版內容改了記得回來同步改這支腳本裡對應的
文字，兩邊沒有自動連動，這是已知的維護成本（見學習筆記.md「靜態站產生器容易跟
批次改版脫節」那條）。

2026-09-11 修正：第一版漏掉「適用機型」那段的官方認證清單超連結
（Steve 發現：Word 檔案裡面沒有適用機型的查詢連結），這版補上真正可點的超連結，
不是只留文字。
"""
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CERT_LIST_URL = "https://support.google.com/chromeosflex/answer/11513094"


def add_hyperlink(paragraph, url, text, size=11, color="1E3A5F", underline=True):
    """python-docx 沒有內建超連結 API，手動組 XML 插入，讓 Word 裡真的點得下去。"""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")

    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    rpr.append(rFonts)

    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(sz)

    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(u)

    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    rpr.append(c)

    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink

ACCENT = RGBColor(0x1E, 0x3A, 0x5F)
DANGER = RGBColor(0xA3, 0x20, 0x17)
WARN = RGBColor(0x8A, 0x5A, 0x00)
SAFE = RGBColor(0x0F, 0x6B, 0x45)
MUTED = RGBColor(0x5C, 0x5C, 0x5C)


def set_cjk_font(run, name="Microsoft JhengHei", size=11):
    run.font.name = name
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name)


def new_doc(title, subtitle):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = "Microsoft JhengHei"
    style.font.size = Pt(11)
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), "Microsoft JhengHei")

    for section in doc.sections:
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)

    p = doc.add_paragraph()
    r = p.add_run("Hunglun Technology · Deployment Guide")
    set_cjk_font(r, size=9)
    r.font.color.rgb = ACCENT
    r.font.bold = True

    h = doc.add_heading(level=0)
    r = h.add_run(title)
    set_cjk_font(r, size=22)
    r.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    h.paragraph_format.space_after = Pt(4)

    p = doc.add_paragraph()
    r = p.add_run(subtitle)
    set_cjk_font(r, size=11)
    r.font.color.rgb = MUTED
    p.paragraph_format.space_after = Pt(18)

    return doc


def add_h2(doc, text):
    h = doc.add_heading(level=1)
    r = h.add_run(text)
    set_cjk_font(r, size=16)
    r.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)


def add_h3(doc, text):
    h = doc.add_heading(level=2)
    r = h.add_run(text)
    set_cjk_font(r, size=13)
    r.font.color.rgb = ACCENT


def add_p(doc, text, bold=False, color=None, size=11):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_cjk_font(r, size=size)
    r.font.bold = bold
    if color:
        r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(8)
    return p


def add_callout(doc, title, body, color):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("● " + title)
    set_cjk_font(r, size=11.5)
    r.font.bold = True
    r.font.color.rgb = color
    add_p(doc, body, size=10.5, color=MUTED)


def add_steps(doc, steps):
    """steps: list of (title, body) tuples, 編號步驟"""
    for i, (t, b) in enumerate(steps, 1):
        p = doc.add_paragraph()
        r = p.add_run(f"{i}. {t}")
        set_cjk_font(r, size=12)
        r.font.bold = True
        r.font.color.rgb = ACCENT
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        add_p(doc, b, size=10.5)


def add_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ''
        r = hdr[i].paragraphs[0].add_run(h)
        set_cjk_font(r, size=10)
        r.font.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            r = cells[i].paragraphs[0].add_run(val)
            set_cjk_font(r, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return t


def linkify_cell(table, row_idx, col_idx, url, text):
    """把表格裡某一格的純文字換成真的可點超連結（1-based 資料列，跳過表頭）。"""
    cell = table.rows[row_idx].cells[col_idx]
    p = cell.paragraphs[0]
    p.clear()
    add_hyperlink(p, url, text, size=10)


def add_link_paragraph(doc, before, link_text, url, after, size=11):
    """一段文字中間夾一個真超連結，格式跟 add_p 一致。"""
    p = doc.add_paragraph()
    if before:
        r = p.add_run(before)
        set_cjk_font(r, size=size)
    add_hyperlink(p, url, link_text, size=size)
    if after:
        r2 = p.add_run(after)
        set_cjk_font(r2, size=size)
    p.paragraph_format.space_after = Pt(8)
    return p


def add_bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style='List Bullet')
        r = p.add_run(it)
        set_cjk_font(r, size=10.5)


# ── 共用內容 ──

WHAT_INTRO = ("Google 官方有一套叫 FRD（ChromeOS Flex Remote Deployment）的工具，是一支 Windows 執行檔 "
              "agent.exe 加上一個設定檔。把它推送到還在跑 Windows 的電腦上執行，它會自動完成四件事：")
WHAT_STEPS = ["清除 TPM、重新分割整顆硬碟", "安裝 ChromeOS Flex", "連上你指定的 Wi-Fi", "用註冊權杖自動加入 Google Admin Console"]

CERT_TABLE_HEAD = ["狀態", "官方定義", "我們的做法"]
CERT_TABLE_ROWS = [
    ["通過認證", "裝置型號預期可與 ChromeOS Flex 搭配運作", "可以接。這是唯一建議部署的狀態"],
    ["預期會出現輕微問題", "可能至少可以支援基本功能，但團隊仍在改善相容性", "先跟客戶說清楚有風險，並實測過再決定"],
    ["預期會出現重大問題", "已知會出現啟動問題等重大問題，目前不建議搭配使用", "不要做"],
    ["取消認證", "已達支援終止日期，支援服務於指定年度 12 月 31 日結束", "不要做"],
]

BEFORE_HEAD = ["要準備的", "哪裡拿", "備註"]

START_STEPS_COMMON = [
    ("雙擊「開始操作.bat」", "瀏覽器會自動打開操作畫面。這支服務只跑在你自己的電腦上，別人連不進來，關掉黑色視窗就結束。"),
    ("輸入系統密碼", "跟公司 CEU 綁定系統同一組密碼。"),
    ("按右上角「準備套件」", "第一次要下載約 1 GB 的官方檔案，記錄區會顯示進度。下載過一次之後就不用再按。"),
    ("填三個欄位並儲存", "註冊權杖、Wi-Fi 名稱、Wi-Fi 密碼。三個沒填完，下面的執行按鈕會是灰的按不下去。"),
    ("確認是綠色的模擬測試模式", "畫面中間那條大橫幅是綠色就對了。這個模式只做檢查，不會動任何硬碟。"),
]

SAFETY_STEPS = [
    ("綠色模式，一到兩台測試機", "確認連得上、設定檔正確、檢查都通過。"),
    ("紅色模式，同樣那一兩台", "真的轉一次，然後去 Admin Console 確認裝置有出現、政策有套用。這步驗的是「整條路走得通」。"),
    ("紅色模式，全部裝置", "前面兩步都乾淨才做這步。"),
]

VERIFY_HEAD = ["要看的地方", "看什麼"]
VERIFY_ROWS_COMMON = [
    ["裝置畫面", "兩分鐘內出現安裝通知並重開機，重開後進入 ChromeOS Flex 的安裝流程"],
    ["Google Admin Console", "裝置 → ChromeOS，新機器有出現、在正確的組織單位、政策有套用。這是最終判準"],
    ["agent.log", "失敗時查原因用。模擬測試會留下 InstallSkipped 記錄，代表檢查通過但依設定沒有真的安裝"],
]

TROUBLE_HEAD = ["狀況", "原因與處理"]
TROUBLE_ROWS_SHARED = [
    ["執行按鈕是灰的", "三個設定欄位沒填完。填好按「儲存設定」就會解鎖"],
    ["裝置轉完了但 Admin Console 沒出現", "兩個可能：註冊權杖填錯，或裝置連不上你設定的 Wi-Fi。先確認 Wi-Fi 名稱與密碼正確"],
    ["「公司鏡像」按不下去", "鏡像還沒建立。用「Google 官方」就好，功能完全一樣"],
    ["模擬測試出現 GetSupportedSize 錯誤", "Windows 的磁碟重組服務設成自動啟動造成的。設定檔裡有一行 disable_defragsvc，取消註解改成 true 即可"],
    ["畫面說已經有工作在跑", "正常的保護機制，同一時間只允許一項工作。等它跑完再按"],
    ["雙擊沒反應／被防毒攔截", "flex-deploy.exe 已內建 Python，不用另外安裝。防毒軟體對沒簽章的執行檔偶爾會誤判，選「仍要執行」或加入白名單即可"],
    ["下載時 Chrome 跳出「這個檔案不是常見的下載項目」", "正常現象，不是抓錯或中毒。zip 裡包了沒有數位簽章的 flex-deploy.exe，又是新建檔案沒有下載記錄，Chrome 才會示警。點「下載可疑的檔案」繼續即可"],
]

APPENDIX_TABLE_HEAD = ["你的情況", "走哪條", "為什麼"]
APPENDIX_TABLE_ROWS = [
    ["Surface Go，S 模式，被 Intune 管", "Intune", "S 模式不准跑 exe，要先用 Intune 政策整批退出 S 模式；而且是 Entra ID 環境，PsExec 沒有帳號可用"],
    ["電腦已經被 Intune 管著", "Intune", "現成的管道，不用另外要帳密。而且多半也是 Entra ID 環境"],
    ["電腦分散在不同校區、你人到不了也連不進內網", "Intune", "Intune 靠裝置自己向雲端報到，裝置在哪都收得到"],
    ["對方資安政策擋掉檔案分享（445 埠）", "Intune", "PsExec 那條路直接斷了"],
    ["傳統網域環境，你在現場同網段，手上有網域管理員帳密", "PsExec", "按下去立刻執行，不用等 Intune 排程"],
    ["就幾台，都在你面前", "PsExec", "最快。設定填一填就跑完了"],
]

GLOSSARY = [
    ("FRD", "ChromeOS Flex Remote Deployment，Google 官方的遠端部署套件，核心是 agent.exe 這支程式。"),
    ("授權池", "組織帳號裡購買好、還沒指派給裝置的授權額度。買授權是先囤在池子裡，裝置註冊時才從池子扣一個名額。"),
    ("註冊權杖", "Enrollment Token。一串代碼，讓裝置轉完之後知道要自動加入哪個單位的 Google Admin Console。等同密碼等級的機密。"),
    ("S 模式", "Windows 的一種鎖定狀態，只准跑 Microsoft Store 的程式，不准跑一般的 .exe。退出是單向的，退了回不去。"),
    ("Win32 應用程式", "就是一般的 Windows 程式（.exe 那種）。agent.exe 屬於這一類，所以在 S 模式下跑不動。"),
    ("Entra ID", "微軟的雲端帳號系統，舊名 Azure AD。電腦加入 Entra ID 後，員工用公司 Microsoft 帳號登入，跟傳統公司網域不同。這種電腦沒有網域管理員帳號，PsExec 用不了。"),
    ("模擬測試", "Dry run。只跑檢查不做實際變更的模式，用來確認流程正確。"),
    ("TPM", "電腦主機板上的安全晶片，存放加密金鑰。轉換系統時要清除它，原本 Windows 的加密資料也就跟著失效。"),
    ("ONC", "Open Network Configuration，ChromeOS 用的網路設定格式。Wi-Fi 的名稱與密碼就是寫在這種檔案裡。"),
]


def add_footer(doc):
    add_p(doc, "")
    p = doc.add_paragraph()
    r = p.add_run("鴻綸科技 Hunglun Technology・Google for Education Professional Development Partner")
    set_cjk_font(r, size=9)
    r.font.color.rgb = MUTED
    p2 = doc.add_paragraph()
    r2 = p2.add_run("Steve · September 2026")
    set_cjk_font(r2, size=9)
    r2.font.color.rgb = MUTED


def build_intune():
    doc = new_doc(
        "ChromeOS Flex 遠端部署手冊：路線 A 必須使用 Intune",
        "裝置是 S 模式、被 Intune／Entra ID 管理、或你人到不了現場，都屬於這條。另一條路線見《ChromeOS Flex 部署手冊：路線 B 可以使用 PsExec》。"
    )

    add_h2(doc, "這條路線適合你嗎")
    add_callout(doc, "符合任一項，就是這條路線", "① 裝置是 S 模式（Surface Go 系列出廠預設就是）② 被 Intune／Entra ID 管理 ③ 你人到不了、也連不進內網 ④ 對方防火牆擋掉檔案分享（445 埠）", WARN)

    add_h2(doc, "這個工具在做什麼")
    add_p(doc, WHAT_INTRO)
    add_bullets(doc, WHAT_STEPS)
    add_p(doc, "在這條 Intune 路線裡，「送過去」的方式是把 agent.exe 包成 Win32 App，交給 Intune 派送，裝置自己向雲端報到、拉下來安裝。")
    add_callout(doc, "這是不可逆的操作", "執行之後，那台電腦原本的 Windows 與裡面所有資料會被清空，變成 ChromeOS Flex，救不回來。整套流程都圍繞「先模擬測試確認，再動真的」設計。", DANGER)

    add_h3(doc, "適用機型：只做官方認證的型號")
    add_callout(doc, "未認證的機器不要碰", "Google 官方明文寫著：裝置如果未列在認證裝置清單中，就不適用 Chrome Education 或 Chrome Enterprise 支援服務。", DANGER)
    add_link_paragraph(doc, "接案前第一件事就是拿對方的機型去對 Google 的", "通過認證的型號清單", CERT_LIST_URL, "。清單上的每個型號有四種狀態，只有第一種可以做：")
    add_table(doc, CERT_TABLE_HEAD, CERT_TABLE_ROWS)
    add_p(doc, "Surface Go 系列特別注意：查證過只有 Surface Go 2、Surface Go 3 是 Certified（支援到 2028 年），第一代 Surface Go 不在清單上，務必核對完整機型名稱。", bold=True)

    add_h2(doc, "開始之前要準備")
    t_before = add_table(doc, BEFORE_HEAD, [
        ["確認機型通過認證", "Google 官方認證型號清單", "接案第一件事。狀態不是「通過認證」不要做"],
        ["確認是不是 S 模式", "裝置上：設定 → 系統 → 啟用", "Surface Go 出廠預設就是。是的話要先用 Intune 政策整批退出"],
        ["Intune 授權", "問對方 IT／採購", "要另外付費，不是每個客戶都有"],
        ["ChromeOS Flex 註冊權杖", "對方單位的 Google Admin Console", "裝置 → Chrome → 註冊權杖 → 註冊，權杖沒有到期日"],
        ["Wi-Fi 名稱與密碼", "現場網管", "裝置轉完之後要靠這組連網才註冊得上"],
        ["官方套件（約 1 GB）", "介面上按「準備套件」自動下載", "出發前先在公司網路下載好"],
        ["一到兩台測試機", "現場挑", "絕對不能跳過"],
    ])
    linkify_cell(t_before, 1, 1, CERT_LIST_URL, "Google 官方認證型號清單")

    add_h2(doc, "三分鐘上手")
    add_steps(doc, START_STEPS_COMMON + [("按「開始產生安裝包」", "系統會產出 .intunewin 檔案，接著到 Intune 後台上傳並指派。")])

    add_h2(doc, "完整步驟")
    add_steps(doc, [
        ("確認機型通過 Google 認證", "對照官方認證清單，精確比對機型名稱與型號。狀態不是「Certified」不要接。"),
        ("檢查是不是 S 模式", "設定 → 系統 → 啟用，版本名稱寫著「S 模式」就是。S 模式的 Windows 不允許安裝執行 Win32 應用程式，agent.exe 正是 Win32 程式，換 PsExec 也沒用。"),
        ("是 S 模式的話，用 Intune 政策整批退出", "有 Intune 授權可以整批遠端退出；沒有的話每台自己到 Microsoft Store 手動退出。這一步是單向的，退出後回不去 S 模式。"),
        ("確認是 Entra ID 環境", "Entra ID 環境的裝置沒有傳統網域管理員帳號可用，PsExec 靠的正是那組帳號。只要是 Entra ID 管理，都只能走 Intune。"),
        ("備妥權杖、Wi-Fi、官方套件", "到對方 Google Admin Console 產生註冊權杖，跟現場網管要 Wi-Fi 名稱密碼，出發前先下載好官方套件。"),
        ("產生安裝包，上傳 Intune", "介面按「開始產生安裝包」，把 agent.exe 打包成 Win32 App，到 Intune 後台上傳、設定、指派。"),
        ("先跑 1～2 台模擬測試", "dry_run 模式只檢查不動硬碟，綠色橫幅代表過了，整批下去之前絕對不能跳過。"),
        ("正式部署與現場流程", "模擬測試通過後才關掉 dry_run 跑正式的。"),
        ("驗收", "照檢查清單逐台過一遍，都通過才算完工交件。"),
    ])

    add_h2(doc, "Intune 後台設定細節")
    add_p(doc, "介面按下「開始產生安裝包」後，會在 output\\ 產出一個約 1 GB 的 .intunewin 檔案。接著到 Intune 後台：")
    add_steps(doc, [
        ("新增 Win32 應用程式", "應用程式 → Windows → 新增 → Win32 應用程式，上傳那個 .intunewin。"),
        ("安裝命令填 agent.exe", "要用系統管理員權限執行。"),
        ("偵測規則填「agent.log 檔案存在」", "這條規則只是滿足 Intune 介面必填要求，真正的成功判斷是回 Admin Console 看。"),
        ("指派對象先選測試群組", "先指派給那一兩台測試機，確認走得通再改成全部裝置。"),
    ])
    add_callout(doc, "Intune 不會立刻執行", "派送時間由 Intune 的同步排程決定，可能要等一段時間裝置才收到，這是正常的。", WARN)

    add_h2(doc, "模擬測試與正式部署")
    add_callout(doc, "綠色：模擬測試模式", "系統只會做檢查，不會動硬碟、不會安裝任何東西，可以放心重複跑。", SAFE)
    add_callout(doc, "紅色：正式部署模式", "真的執行轉換，裝置的 Windows 與所有資料會被清空，沒有還原的方法。", DANGER)
    add_p(doc, "切換方式：按橫幅右邊的「切換到正式部署」，必須手動輸入「確認部署」四個字才切得過去。")
    add_h3(doc, "正確的順序")
    add_steps(doc, SAFETY_STEPS)

    add_h2(doc, "現場完整流程")
    add_steps(doc, [
        ("接案時：先查機型有沒有通過認證", "狀態不是「通過認證」就不要接。"),
        ("接案時：確認是不是 S 模式、有沒有 Intune 授權", "到現場才發現會很難處理。"),
        ("出發前：在公司先按一次「準備套件」", "把 1 GB 的官方檔案抓下來。"),
        ("到現場：跟對方索取註冊權杖", "請對方的 Google 管理員到 Admin Console 產生。"),
        ("填三個欄位並儲存", "註冊權杖、Wi-Fi 名稱、Wi-Fi 密碼，保持綠色模擬測試模式。"),
        ("對一兩台測試機跑模擬測試", "看記錄區有沒有錯誤訊息，有錯先解決。"),
        ("切成紅色，同樣那一兩台跑真的", "裝置會在兩分鐘內顯示安裝通知並重新開機。"),
        ("去 Admin Console 確認註冊成功", "新機器有沒有出現在正確的組織單位、政策有沒有套用。"),
        ("全部沒問題，才把指派對象改成全部裝置", "Intune 後台把指派群組從測試機擴大到全部目標裝置。"),
    ])

    add_h2(doc, "怎麼確認成功")
    add_p(doc, "Google 官方明確建議：用 Admin Console 或 Directory API 確認裝置有沒有註冊進來，不要只看本機記錄。")
    add_table(doc, VERIFY_HEAD, VERIFY_ROWS_COMMON)

    add_h2(doc, "出問題怎麼辦")
    add_table(doc, TROUBLE_HEAD, TROUBLE_ROWS_SHARED + [
        ["Intune 派送很久都沒動靜", "正常，派送時間由 Intune 同步排程決定。急著要立刻執行只能改走路線 B：PsExec（前提是裝置符合那條路線的條件）"],
    ])

    add_h2(doc, "給技術人員：指令列")
    add_p(doc, "設定檔在 secrets\\config.json，欄位跟介面上的一樣。")
    add_h3(doc, "產生 Intune 安裝包")
    add_p(doc, "python intune_flex_packager.py")
    add_h3(doc, "建立公司鏡像")
    add_p(doc, "python mirror_to_r2.py --check   # 先驗憑證\npython mirror_to_r2.py           # 實際上傳")

    add_h2(doc, "名詞解釋")
    for term, defn in GLOSSARY:
        p = doc.add_paragraph()
        r = p.add_run(term + "：")
        set_cjk_font(r, size=10.5)
        r.font.bold = True
        r.font.color.rgb = ACCENT
        r2 = p.add_run(defn)
        set_cjk_font(r2, size=10.5)

    add_h2(doc, "附錄：判斷邏輯")
    add_p(doc, "想弄懂「為什麼是這條路線」背後的完整推理，看這節。")
    add_h3(doc, "問題一：機型有沒有通過 Google 認證？")
    add_p(doc, "沒有就不做，這案子不該接。")
    add_h3(doc, "問題二：電腦是不是「S 模式」？")
    add_callout(doc, "S 模式的電腦，兩條路線都會失敗", "微軟官方寫得很清楚：Windows S 模式裝置預設不允許安裝和執行 Win32 應用程式。agent.exe 就是 Win32 程式，檔案送得過去但跑不動，換 PsExec 也沒用。", DANGER)
    add_h3(doc, "問題三：這些電腦是被 Intune（Entra ID）管的嗎？")
    add_callout(doc, "這種電腦，PsExec 實務上用不了", "只加入 Entra ID 的機器沒有傳統網域管理員帳號可用，PsExec 靠的正是那組帳號，所以只剩 Intune 一條路。", WARN)
    add_h3(doc, "問題四：你的電腦連得到那些機器嗎？")
    add_p(doc, "前面三題都過了才輪到這題。連得到就可以用 PsExec，連不到就走 Intune。")
    add_h3(doc, "結論：一張表看完")
    add_table(doc, APPENDIX_TABLE_HEAD, APPENDIX_TABLE_ROWS)

    add_footer(doc)
    return doc


def build_psexec():
    doc = new_doc(
        "ChromeOS Flex 遠端部署手冊：路線 B 可以使用 PsExec",
        "要同時符合：不是 S 模式、傳統網域環境、人在現場同網段、有網域管理員帳密。另一條路線見《ChromeOS Flex 部署手冊：路線 A 必須使用 Intune》。"
    )

    add_h2(doc, "這條路線適合你嗎")
    add_callout(doc, "四個條件要全部符合", "① 裝置不是 S 模式 ② 是傳統網域環境，不是 Entra ID Only ③ 你人在現場、跟裝置同一個網段連得到 ④ 手上有網域管理員帳密。缺一項就要改走路線 A：必須用 Intune。", SAFE)

    add_h2(doc, "這個工具在做什麼")
    add_p(doc, WHAT_INTRO)
    add_bullets(doc, WHAT_STEPS)
    add_p(doc, "在這條 PsExec 路線裡，「送過去」的方式是用網域管理員帳密直接連進每一台電腦，把 agent.exe 推過去並立刻執行。")
    add_callout(doc, "這是不可逆的操作", "執行之後，那台電腦原本的 Windows 與裡面所有資料會被清空，變成 ChromeOS Flex，救不回來。整套流程都圍繞「先模擬測試確認，再動真的」設計。", DANGER)

    add_h3(doc, "適用機型：只做官方認證的型號")
    add_callout(doc, "未認證的機器不要碰", "Google 官方明文寫著：裝置如果未列在認證裝置清單中，就不適用 Chrome Education 或 Chrome Enterprise 支援服務。", DANGER)
    add_link_paragraph(doc, "接案前第一件事就是拿對方的機型去對 Google 的", "通過認證的型號清單", CERT_LIST_URL, "。清單上的每個型號有四種狀態，只有第一種可以做：")
    add_table(doc, CERT_TABLE_HEAD, CERT_TABLE_ROWS)

    add_h2(doc, "開始之前要準備")
    t_before = add_table(doc, BEFORE_HEAD, [
        ["確認機型通過認證", "Google 官方認證型號清單", "接案第一件事。狀態不是「通過認證」不要做"],
        ["確認不是 S 模式", "裝置上：設定 → 系統 → 啟用", "是 S 模式的話 PsExec 用不了，要改走路線 A"],
        ["網域管理員帳密", "對方 IT", "那個帳號要在所有目標裝置上都有系統管理員權限"],
        ["ChromeOS Flex 註冊權杖", "對方單位的 Google Admin Console", "裝置 → Chrome → 註冊權杖 → 註冊，權杖沒有到期日"],
        ["Wi-Fi 名稱與密碼", "現場網管", "裝置轉完之後要靠這組連網才註冊得上"],
        ["官方套件（約 1 GB）", "介面上按「準備套件」自動下載", "出發前先在公司網路下載好"],
        ["一到兩台測試機", "現場挑", "絕對不能跳過"],
    ])
    linkify_cell(t_before, 1, 1, CERT_LIST_URL, "Google 官方認證型號清單")

    add_h2(doc, "三分鐘上手")
    add_steps(doc, START_STEPS_COMMON + [("切到 PsExec 分頁，填裝置清單並執行", "填好目標電腦清單，按「開始部署」。")])

    add_h2(doc, "完整步驟")
    add_steps(doc, [
        ("確認機型通過 Google 認證", "對照官方認證清單，精確比對機型名稱與型號。狀態不是「Certified」不要接。"),
        ("確認不是 S 模式", "設定 → 系統 → 啟用，確認版本名稱沒有寫「S 模式」。是的話 PsExec 用不了，要改走路線 A 先退出 S 模式。"),
        ("確認是傳統網域、連得到、有帳密", "員工登入用的是傳統公司網域帳號；你人在現場、跟這批電腦同一個網段；手上有網域管理員帳密。三項缺一都走不下去。"),
        ("備妥權杖、Wi-Fi、官方套件", "到對方 Google Admin Console 產生註冊權杖，跟現場網管要 Wi-Fi 名稱密碼，出發前先下載好官方套件。"),
        ("填裝置清單，執行 PsExec 部署", "介面用網域管理員帳密透過 PsExec 直接連進每一台電腦、推送並執行 agent.exe，按下去立刻跑。"),
        ("先跑 1～2 台模擬測試", "dry_run 模式只檢查不動硬碟，綠色橫幅代表過了，整批下去之前絕對不能跳過。"),
        ("正式部署與現場流程", "模擬測試通過後才關掉 dry_run 跑正式的。"),
        ("驗收", "照檢查清單逐台過一遍，都通過才算完工交件。"),
    ])

    add_h2(doc, "PsExec 操作細節")
    add_h3(doc, "裝置清單怎麼填")
    add_p(doc, "表格一行一台，填電腦名稱或 IP。可以直接從 Excel 整欄複製後貼上，系統會自動展開成多列。")
    add_h3(doc, "帳號密碼")
    add_p(doc, "兩個欄位留空的話，會用你目前登入這台電腦的帳號權限去連。不是的話就填一組有權限的網域管理員帳密。")
    add_h3(doc, "同時處理幾台")
    add_p(doc, "預設 5 台，數字調大速度快但網路負擔重，一般不用超過 10。")
    add_h3(doc, "系統實際做的事")
    add_p(doc, "對清單上每一台裝置依序：連上管理共用資料夾、複製 FRD 套件、用 PsExec 遠端執行 agent.exe、斷開連線。任何一步失敗不影響其他台繼續跑。")
    add_h3(doc, "跑完的報告")
    add_p(doc, "畫面即時顯示成功與失敗台數，全部跑完可下載 CSV 報告，用 Excel 打開查看每一台的詳細結果。")
    add_callout(doc, "連不上通常是防火牆", "PsExec 需要對方裝置開放檔案分享（445 埠）與管理共用資料夾。先拿一台測試確認通得過，再排整批。", WARN)

    add_h2(doc, "模擬測試與正式部署")
    add_callout(doc, "綠色：模擬測試模式", "系統只會做檢查，不會動硬碟、不會安裝任何東西，可以放心重複跑。", SAFE)
    add_callout(doc, "紅色：正式部署模式", "真的執行轉換，裝置的 Windows 與所有資料會被清空，沒有還原的方法。", DANGER)
    add_p(doc, "切換方式：按橫幅右邊的「切換到正式部署」，必須手動輸入「確認部署」四個字才切得過去。PsExec 還有第二道保護：按下「開始部署」會再問一次「確定要對 N 台裝置執行嗎」。")
    add_h3(doc, "正確的順序")
    add_steps(doc, SAFETY_STEPS)

    add_h2(doc, "現場完整流程")
    add_steps(doc, [
        ("接案時：先查機型有沒有通過認證", "狀態不是「通過認證」就不要接。"),
        ("接案時：確認不是 S 模式、確認網域環境", "問清楚是不是傳統網域環境、你到時候連不連得到內網。"),
        ("出發前：在公司先按一次「準備套件」", "把 1 GB 的官方檔案抓下來。"),
        ("到現場：跟對方索取註冊權杖", "請對方的 Google 管理員到 Admin Console 產生。"),
        ("填三個欄位並儲存", "註冊權杖、Wi-Fi 名稱、Wi-Fi 密碼，保持綠色模擬測試模式。"),
        ("對一兩台測試機跑模擬測試", "看記錄區有沒有錯誤訊息，有錯先解決。"),
        ("切成紅色，同樣那一兩台跑真的", "裝置會在兩分鐘內顯示安裝通知並重新開機。"),
        ("去 Admin Console 確認註冊成功", "新機器有沒有出現在正確的組織單位、政策有沒有套用。"),
        ("全部沒問題，才把裝置清單填完整", "把 PsExec 的裝置清單從測試機擴大到全部目標裝置。"),
    ])

    add_h2(doc, "怎麼確認成功")
    add_p(doc, "Google 官方明確建議：用 Admin Console 或 Directory API 確認裝置有沒有註冊進來，不要只看本機記錄。")
    add_table(doc, VERIFY_HEAD, VERIFY_ROWS_COMMON + [["PsExec 跑完的 CSV 報告", "先看有沒有回報失敗的台數，失敗的個別去查 agent.log"]])

    add_h2(doc, "出問題怎麼辦")
    add_table(doc, TROUBLE_HEAD, [
        ["執行按鈕是灰的", "三個設定欄位沒填完。填好按「儲存設定」就會解鎖"],
        ["System error 67：The network name cannot be found", "PsExec 找不到那台電腦。確認電腦名稱或 IP 打對、機器開著、跟你在同一個網路"],
        ["檔案送過去了，但裝置上什麼都沒發生", "先確認那台是不是 S 模式。S 模式不准跑 .exe，這種情況要改走路線 A"],
        ["連得到但複製檔案失敗", "多半是權限或防火牆。確認帳號在對方機器上是系統管理員，以及 445 埠沒被擋"],
    ] + TROUBLE_ROWS_SHARED[1:])

    add_h2(doc, "給技術人員：指令列")
    add_p(doc, "設定檔在 secrets\\config.json，欄位跟介面上的一樣。")
    add_h3(doc, "PsExec 大量部署")
    add_p(doc, "python psexec_mass_deploy.py --parallel 5")
    add_p(doc, "裝置清單讀 secrets\\devices.txt，一行一台，# 開頭是註解。dry_run 為 false 時必須加 --yes 才會真的執行。")
    add_h3(doc, "建立公司鏡像")
    add_p(doc, "python mirror_to_r2.py --check   # 先驗憑證\npython mirror_to_r2.py           # 實際上傳")

    add_h2(doc, "名詞解釋")
    for term, defn in GLOSSARY:
        p = doc.add_paragraph()
        r = p.add_run(term + "：")
        set_cjk_font(r, size=10.5)
        r.font.bold = True
        r.font.color.rgb = ACCENT
        r2 = p.add_run(defn)
        set_cjk_font(r2, size=10.5)

    add_h2(doc, "附錄：判斷邏輯")
    add_p(doc, "想弄懂「為什麼是這條路線」背後的完整推理，看這節。")
    add_h3(doc, "問題一：機型有沒有通過 Google 認證？")
    add_p(doc, "沒有就不做，這案子不該接。")
    add_h3(doc, "問題二：電腦是不是「S 模式」？")
    add_callout(doc, "S 模式的電腦，兩條路線都會失敗", "微軟官方寫得很清楚：Windows S 模式裝置預設不允許安裝和執行 Win32 應用程式。agent.exe 就是 Win32 程式，檔案送得過去但跑不動，換 PsExec 也沒用。", DANGER)
    add_h3(doc, "問題三：這些電腦是被 Intune（Entra ID）管的嗎？")
    add_callout(doc, "這種電腦，PsExec 實務上用不了", "只加入 Entra ID 的機器沒有傳統網域管理員帳號可用，PsExec 靠的正是那組帳號，所以只剩 Intune 一條路。", WARN)
    add_h3(doc, "問題四：你的電腦連得到那些機器嗎？")
    add_p(doc, "前面三題都過了才輪到這題。連得到就可以用 PsExec，連不到就走 Intune。")
    add_h3(doc, "結論：一張表看完")
    add_table(doc, APPENDIX_TABLE_HEAD, APPENDIX_TABLE_ROWS)

    add_footer(doc)
    return doc


if __name__ == "__main__":
    out_dir = sys.argv[1] if len(sys.argv) > 1 else "."
    d1 = build_intune()
    d1.save(f"{out_dir}/路線A-Intune部署手冊.docx")
    d2 = build_psexec()
    d2.save(f"{out_dir}/路線B-PsExec部署手冊.docx")
    print("完成")
